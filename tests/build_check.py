"""Check that Publish refuses a broken Sheet and leaves the live questions untouched.
Run: python tests/build_check.py
"""
import pathlib, shutil, subprocess, sys, tempfile
from openpyxl import load_workbook
sys.stdout.reconfigure(encoding="utf-8")
ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "tools" / "1MV_Listening_Form_V2_Questions.xlsx"
LIVE = ROOT / "content" / "questions.json"
before = LIVE.read_bytes()
tmp = pathlib.Path(tempfile.mkdtemp()) / "broken.xlsx"
shutil.copy(SRC, tmp)
wb = load_workbook(tmp)
q = wb["Questions"]; head = [c.value for c in q[1]]
col = {h: i + 1 for i, h in enumerate(head)}
for r in range(2, q.max_row + 1):
    qid = q.cell(r, col["id"]).value
    if qid == "Y11a_work": q.cell(r, col["question_en"]).value = None          # missing English question
    if qid == "Y6_class": q.cell(r, col["show_if"]).value = "Y5 in school, univ"  # unknown option code
    if qid == "Y7_institution": q.cell(r, col["options"]).value = "nolist"      # missing options list
    if qid == "G1_age": q.cell(r, col["type"]).value = "slider"                # unknown type
wb.save(tmp)
r = subprocess.run([sys.executable, "-I", str(ROOT / "tools" / "build_questions.py"), "--xlsx", str(tmp)], capture_output=True, text=True, encoding="utf-8")
out = r.stdout
checks = [
    (r.returncode == 1, "broken Sheet → Publish stops (exit 1)"),
    ("Y11a_work: question_en is empty" in out, "reports the missing English question"),
    ("'univ' is not an option code of Y5" in out, "reports the wrong show_if code"),
    ("options list 'nolist' has no rows" in out, "reports the missing options list"),
    ("type 'slider'" in out, "reports the unknown type"),
    (LIVE.read_bytes() == before, "live questions.json left unchanged"),
]
ok = r2 = subprocess.run([sys.executable, "-I", str(ROOT / "tools" / "build_questions.py"), "--xlsx", str(SRC)], capture_output=True, text=True, encoding="utf-8")
checks.append((r2.returncode == 0, "the good Sheet still builds"))
LIVE.write_bytes(before)  # put the published file (with its audio list) back exactly as it was
for good, msg in checks:
    print(("ok   " if good else "FAIL ") + msg)
fails = sum(not g for g, _ in checks)
print("all passed" if not fails else f"{fails} FAILED")
sys.exit(1 if fails else 0)
