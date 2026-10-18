"""Build the "Questions (source)" workbook from tools/sheet_seed.json (made by export_to_sheet.mjs).
Run: python tools/make_seed_xlsx.py  ->  tools/1MV_Listening_Form_V2_Questions.xlsx (upload to Google Drive as a Google Sheet)
"""
import json, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).parent
d = json.loads((HERE / "sheet_seed.json").read_text(encoding="utf-8"))
LANG_FILL = {"en": "DCEBFA", "hi": "FDF1DC", "mr": "E2F2E7"}
HEAD = PatternFill("solid", fgColor="1F5F8B")
LOCK = PatternFill("solid", fgColor="EEF2F7")

wb = Workbook()
ws = wb.active
ws.title = "READ ME"
readme = [
    ("1MV Listening Form V2 — Questions (source)", True),
    ("This Sheet is where the form's questions come from. Edit here, then press Publish (GitHub → Actions → Publish questions → Run workflow). The form updates in 2–3 minutes, with fresh audio.", False),
    ("", False),
    ("Golden rules", True),
    ("1. Never change an ID (column 'id') or an option 'code'. Change the words as much as you like.", False),
    ("2. Grey columns are structure (form, id, type, rules). Coloured columns are words in each language — translators edit only these.", False),
    ("3. An empty translation shows the English text instead, so nothing breaks.", False),
    ("4. Every short-answer question asks at most ONE follow-up. Its words are in followup_<language>. Leave the 'followup' column empty to switch it off.", False),
    ("5. Publish checks everything first. If something is wrong (e.g. a missing option), Publish stops, says what to fix, and the live form keeps running unchanged.", False),
    ("", False),
    ("Add a language (e.g. Kannada, code kn)", True),
    ("a. Languages tab: add a row: code kn, name ಕನ್ನಡ, speech_code kn-IN, voice kn-IN-SapnaNeural, enabled yes.", False),
    ("b. Add columns with that code: Questions → question_kn, example_kn, followup_kn · Options → label_kn · Forms → title_kn, intro_kn · Sections → title_kn, note_kn · Screen text → kn.", False),
    ("c. Add common non-answers in that language (e.g. 'yes', 'don't know') to Answer checks → non_answer.", False),
    ("", False),
    ("Tabs", True),
    ("Languages — which languages are on, and the voice used for 🔊 audio.", False),
    ("Forms — title and intro of the Youth, Parent and Public forms.", False),
    ("Sections — section names and notes.", False),
    ("Questions — one row per question, in the order asked. show_if makes a question appear only when another answer fits, e.g. 'Y5 in school, college' or 'Y14_now = notyet'.", False),
    ("Options — answer choices. A question uses the list named in its 'options' column (its own ID, or a shared list like @states).", False),
    ("General words — words that alone do not answer a question (e.g. good, peaceful, hard work). If an answer has only these, the follow-up is asked.", False),
    ("Answer checks — non-answers, 'no one' phrases, child's-choice phrases, safety (distress) words, and filler words.", False),
    ("Screen text — buttons and messages.", False),
]
for i, (text, bold) in enumerate(readme, 1):
    c = ws.cell(i, 1, text)
    c.font = Font(bold=bold, size=14 if i == 1 else 11, color="1F5F8B" if bold else "1B2230")
    c.alignment = Alignment(wrap_text=True, vertical="top")
ws.column_dimensions["A"].width = 130

def sheet(name, rows, widths=None, choices=None):
    ws = wb.create_sheet(name)
    cols = list(rows[0].keys())
    for j, k in enumerate(cols, 1):
        c = ws.cell(1, j, k)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for i, r in enumerate(rows, 2):
        for j, k in enumerate(cols, 1):
            v = r.get(k, "")
            c = ws.cell(i, j, "" if v is None else v)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            lang = k.rsplit("_", 1)[-1] if "_" in k else k
            if lang in LANG_FILL and (k.startswith(("question_", "example_", "followup_", "label_", "title_", "intro_", "note_")) or k in LANG_FILL):
                c.fill = PatternFill("solid", fgColor=LANG_FILL[lang])
            else:
                c.fill = LOCK
    for j, k in enumerate(cols, 1):
        w = (widths or {}).get(k)
        if w is None:
            w = 42 if k.startswith(("question_", "example_", "followup_", "intro_", "note_")) or k in ("words", "meaning") else 26 if k.startswith(("label_", "title_")) or k in LANG_FILL else 14
        ws.column_dimensions[ws.cell(1, j).column_letter].width = w
    ws.freeze_panes = "C2" if "id" in cols else "B2"
    for col, values in (choices or {}).items():
        j = cols.index(col) + 1
        letter = ws.cell(1, j).column_letter
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        ws.add_data_validation(dv); dv.add(f"{letter}2:{letter}1000")
    return ws

sheet("Languages", d["languages"], {"name": 14, "voice": 24})
sheet("Forms", d["forms"])
sheet("Sections", d["sections"])
sheet("Questions", d["questions"], {"show_if": 26, "options": 14, "followup_general_words": 18, "followup_accepts": 16},
      {"type": ["short", "para", "number", "choice", "multi", "dropdown"], "required": ["yes"], "followup": ["yes"], "other_box": ["yes"],
       "followup_accepts": ["no one", "child's choice"]})
sheet("Options", d["options"], {"list": 16, "code": 14, "order": 8})
sheet("General words", d["general_words"], {"list": 20, "words": 120})
sheet("Answer checks", d["answer_checks"], {"check": 16, "meaning": 44, "words": 110})
sheet("Screen text", d["ui"], {"key": 18})
out = HERE / "1MV_Listening_Form_V2_Questions.xlsx"
wb.save(out)
print(out, out.stat().st_size, "bytes")
