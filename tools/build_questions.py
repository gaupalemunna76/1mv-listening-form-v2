"""Build content/questions.json from the "Questions (source)" Google Sheet, after checking everything.

  python tools/build_questions.py --sheet-id <GOOGLE_SHEET_ID>      (what the Publish button runs)
  python tools/build_questions.py --xlsx tools/1MV_Listening_Form_V2_Questions.xlsx   (local copy)

The Sheet must be viewable by anyone with the link (questions are not private). Each tab is read as CSV.
If anything is wrong the script prints every problem and exits with an error, and nothing is written,
so the live form keeps running on the last good questions.
"""
import argparse, csv, hashlib, io, json, os, pathlib, re, sys, urllib.parse, urllib.request

sys.stdout.reconfigure(encoding="utf-8")
ROOT = pathlib.Path(__file__).resolve().parent.parent
OUT = ROOT / "content" / "questions.json"
TABS = ["Languages", "Forms", "Sections", "Questions", "Options", "General words", "Answer checks", "Screen text"]
TYPES = {"short", "para", "number", "choice", "multi", "dropdown"}
CHOICE_TYPES = {"choice", "multi", "dropdown"}
problems, notes = [], []


def read_sheet(sheet_id):
    tabs = {}
    for tab in TABS:
        url = f"https://docs.google.com/spreadsheets/d/{sheet_id}/gviz/tq?tqx=out:csv&sheet={urllib.parse.quote(tab)}"
        try:
            with urllib.request.urlopen(url, timeout=60) as r:
                text = r.read().decode("utf-8")
        except Exception as e:
            sys.exit(f"Could not read tab '{tab}' ({e}). Is the Sheet shared as 'Anyone with the link: Viewer'?")
        if text.lstrip().startswith("<"):
            sys.exit(f"Could not read tab '{tab}': Google returned a web page. Share the Sheet as 'Anyone with the link: Viewer'.")
        tabs[tab] = list(csv.DictReader(io.StringIO(text)))
    return tabs


def read_xlsx(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    tabs = {}
    for tab in TABS:
        if tab not in wb.sheetnames:
            sys.exit(f"Tab '{tab}' is missing.")
        rows = list(wb[tab].iter_rows(values_only=True))
        head = [str(h or "").strip() for h in rows[0]]
        tabs[tab] = [{head[i]: ("" if v is None else str(v)) for i, v in enumerate(r) if i < len(head)} for r in rows[1:]]
    return tabs


def clean(rows):
    out = []
    for r in rows:
        r = {str(k).strip(): ("" if v is None else str(v)).strip() for k, v in r.items() if k is not None}
        if any(r.values()):
            out.append(r)
    return out


def num(v, where):
    if v == "":
        return None
    try:
        f = float(v)
        return int(f) if f.is_integer() else f
    except ValueError:
        problems.append(f"{where}: '{v}' should be a number")


def words(v):
    return [w.strip() for w in re.split(r"[,\n]", v or "") if w.strip()]


def parse_condition(text, where, known):
    """'Y5 in school, college' | 'Y14_now = notyet' | 'P7_visits != 0'  ->  {q, op, values}"""
    m = re.fullmatch(r"\s*([A-Za-z0-9_]+)\s*(=|!=|in|not in)\s*(.+?)\s*", text or "")
    if not m:
        problems.append(f"{where}: show_if '{text}' — write it like 'Y5 in school, college' or 'Y14_now = notyet'")
        return None
    q, op, vals = m.group(1), m.group(2), words(m.group(3))
    if q not in known:
        problems.append(f"{where}: show_if uses '{q}', which is not an earlier question in this form")
        return None
    codes = known[q]
    for v in vals:
        if codes is not None and v not in codes:
            problems.append(f"{where}: show_if value '{v}' is not an option code of {q} (codes: {', '.join(sorted(codes))})")
    return {"q": q, "op": {"=": "in", "in": "in", "!=": "notin", "not in": "notin"}[op], "values": vals}


def build(tabs):
    t = {k: clean(v) for k, v in tabs.items()}
    langs = [l for l in t["Languages"] if l.get("enabled", "").lower() in ("yes", "y", "true", "1")]
    if not langs or langs[0].get("code") != "en":
        problems.append("Languages: English (code en) must be the first enabled language; other languages fall back to it")
    codes = [l["code"] for l in langs]

    def per(row, prefix, where, required_en=False):
        d = {}
        for c in codes:
            v = row.get(f"{prefix}_{c}", "")
            if v:
                d[c] = v
        if required_en and not d.get("en"):
            problems.append(f"{where}: {prefix}_en is empty")
        missing = [c for c in codes if c != "en" and not d.get(c)]
        if d.get("en") and missing:
            notes.append(f"{where}: no {prefix} in {', '.join(missing)} yet (English will show)")
        return d

    # Options lists
    lists = {}
    for i, o in enumerate(t["Options"], 2):
        lst, code = o.get("list", ""), o.get("code", "")
        where = f"Options row {i}"
        if not lst or not code:
            problems.append(f"{where}: list and code are required"); continue
        if code in lists.setdefault(lst, {}):
            problems.append(f"{where}: code '{code}' appears twice in list {lst}")
        lists[lst][code] = {"v": code, "order": num(o.get("order", ""), where) or 0, **per(o, "label", f"{where} ({lst}/{code})", True)}
    general = {g["list"]: words(g.get("words", "")) for g in t["General words"] if g.get("list")}
    checks = {c["check"]: words(c.get("words", "")) for c in t["Answer checks"] if c.get("check")}
    for need in ("non_answer", "no_one", "childs_choice", "distress", "filler"):
        if need not in checks:
            problems.append(f"Answer checks: row '{need}' is missing")

    forms = {}
    for f in t["Forms"]:
        forms[f["form"]] = {"key": f["form"], "title": per(f, "title", f"Forms/{f['form']}", True), "intro": per(f, "intro", f"Forms/{f['form']}", True), "sections": []}
    secs = {}
    for s in sorted(t["Sections"], key=lambda s: (s.get("form", ""), num(s.get("order", ""), "Sections") or 0)):
        if s.get("form") not in forms:
            problems.append(f"Sections: form '{s.get('form')}' is not in the Forms tab"); continue
        sec = {"id": s["section"], "title": per(s, "title", f"Sections/{s['form']}/{s['section']}", True), "note": per(s, "note", f"Sections/{s['form']}/{s['section']}"), "items": []}
        forms[s["form"]]["sections"].append(sec); secs[(s["form"], s["section"])] = sec

    known = {}  # form -> {qid: option codes or None}
    for i, q in enumerate(sorted(t["Questions"], key=lambda q: (q.get("form", ""), num(q.get("order", ""), "Questions") or 0)), 2):
        form, qid, typ = q.get("form", ""), q.get("id", ""), q.get("type", "")
        where = f"Questions {form}/{qid or '(no id)'}"
        if form not in forms:
            problems.append(f"{where}: form '{form}' is not in the Forms tab"); continue
        if not re.fullmatch(r"[A-Za-z][A-Za-z0-9_]*", qid):
            problems.append(f"{where}: id must be letters, numbers and _ only"); continue
        kf = known.setdefault(form, {})
        if qid in kf:
            problems.append(f"{where}: id used twice in this form"); continue
        if typ not in TYPES:
            problems.append(f"{where}: type '{typ}' must be one of {', '.join(sorted(TYPES))}"); continue
        sec = secs.get((form, q.get("section", "")))
        if not sec:
            problems.append(f"{where}: section '{q.get('section')}' is not in the Sections tab for {form}"); continue
        item = {"id": qid, "type": typ, "label": per(q, "question", where, True)}
        ex = per(q, "example", where)
        if ex: item["hint"] = ex
        if q.get("required", "").lower() in ("yes", "y", "true", "1"): item["required"] = True
        if typ == "number":
            item["min"], item["max"] = num(q.get("number_min", ""), where), num(q.get("number_max", ""), where)
            if item["min"] is None or item["max"] is None:
                problems.append(f"{where}: number questions need number_min and number_max")
        codes_here = None
        if typ in CHOICE_TYPES:
            lst = q.get("options", "") or qid
            if lst not in lists:
                problems.append(f"{where}: options list '{lst}' has no rows in the Options tab")
            else:
                item["options"] = [{k: v for k, v in o.items() if k != "order"} for o in sorted(lists[lst].values(), key=lambda o: o["order"])]
                codes_here = set(lists[lst])
            ex_codes = words(q.get("exclusive", ""))
            if ex_codes:
                item["exclusive"] = ex_codes
                for c in ex_codes:
                    if codes_here is not None and c not in codes_here:
                        problems.append(f"{where}: exclusive code '{c}' is not an option")
            if q.get("other_box", "").lower() in ("yes", "y", "true", "1"): item["other"] = True
        if q.get("show_if"):
            cond = parse_condition(q["show_if"], where, kf)
            if cond: item["showIf"] = cond
        if q.get("followup", "").lower() in ("yes", "y", "true", "1"):
            if typ not in ("short", "para"):
                problems.append(f"{where}: follow-ups work only on short/para questions")
            pr = {"minWords": num(q.get("followup_min_words", ""), where) or 3, "ask": per(q, "followup", where, True)}
            gw = q.get("followup_general_words", "")
            if gw:
                if gw not in general:
                    problems.append(f"{where}: general words list '{gw}' is not in the General words tab")
                else:
                    pr["vague"] = general[gw]
            acc = q.get("followup_accepts", "").lower()
            if acc.startswith("no one"): pr["allowNoOne"] = True
            elif acc.startswith("child"): pr["allowChoice"] = True
            elif acc: problems.append(f"{where}: followup_accepts must be 'no one', \"child's choice\" or empty")
            item["probe"] = pr
        sec["items"].append(item)
        kf[qid] = codes_here
    for f in forms.values():
        f["sections"] = [s for s in f["sections"] if s["items"]]
        if not f["sections"]:
            problems.append(f"Forms/{f['key']}: has no questions")
    ui = {}
    for r in t["Screen text"]:
        if r.get("key"):
            ui[r["key"]] = {c: r[c] for c in codes if r.get(c)}
    languages = [{"code": l["code"], "name": l.get("name") or l["code"], "speech": l.get("speech_code") or l["code"], "voice": l.get("voice", "")} for l in langs]
    content = {"languages": languages, "forms": forms, "ui": ui, "checks": checks}
    content["version"] = hashlib.sha1(json.dumps(content, sort_keys=True, ensure_ascii=False).encode()).hexdigest()[:10]
    return content


def main():
    ap = argparse.ArgumentParser()
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--sheet-id")
    g.add_argument("--xlsx")
    ap.add_argument("--report", help="write a Markdown report here (GitHub step summary)")
    a = ap.parse_args()
    tabs = read_sheet(a.sheet_id) if a.sheet_id else read_xlsx(a.xlsx)
    content = build(tabs)
    lines = []
    if problems:
        lines += ["## ❌ Publish stopped — please fix these in the Questions Sheet", ""] + [f"- {p}" for p in problems]
        lines += ["", "The live form was not changed."]
    else:
        n = sum(len(s["items"]) for f in content["forms"].values() for s in f["sections"])
        lines += [f"## ✅ Questions checked: {len(content['forms'])} forms, {n} questions, languages: {', '.join(l['code'] for l in content['languages'])}"]
        OUT.parent.mkdir(exist_ok=True)
        OUT.write_text(json.dumps(content, ensure_ascii=False, indent=1), encoding="utf-8")
        lines += [f"Written content/questions.json (version {content['version']})."]
    if notes:
        lines += ["", "### Notes (not errors)", ""] + [f"- {n}" for n in notes[:200]]
    report = "\n".join(lines)
    print(report)
    if a.report:
        pathlib.Path(a.report).write_text(report + "\n", encoding="utf-8")
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
